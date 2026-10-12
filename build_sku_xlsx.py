# -*- coding: utf-8 -*-
"""agg.json から商品実績(モデル・カラー・サイズ)の Excel を書き出す。

    python build_sku_xlsx.py [出力先.xlsx]

ダッシュボードの生成には使わない補助スクリプト。唯一ここだけ openpyxl を使う
(Excel 側の書式・条件付き書式を素手で組むと壊れやすいため)。

売上は2本立て:
  list = 明細行の金額合計。SiteGiant の明細は定価が入っていて値引きが乗らない
  net  = 注文合計を明細に按分したもの。全体の純売上に一致する
"""
import collections
import datetime
import io
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, u'商品別実績.xlsx')

HEAD_FILL = PatternFill('solid', fgColor='2B2F3A')
HEAD_FONT = Font(bold=True, color='FFFFFF')
TOT_FILL = PatternFill('solid', fgColor='E6ECF6')
THIN = Side(style='thin', color='C9CDD6')
UNDER = Border(bottom=THIN)
MONEY = u'#,##0.00'
INTF = u'#,##0'
PCTF = u'0.0%'
DATEF = u'yyyy-mm-dd'

CATNAME = {'shoes': u'シューズ', 'socks': u'ソックス'}
GRPNAME = {'W': u'レディース(シューズ)', 'M': u'メンズ(シューズ)',
           'tabi': u'足袋(cm)', 'sock': u'ソックス'}
GRPORD = ('W', 'M', 'tabi', 'sock')
SOCKORD = {'XS': 1, 'S': 2, 'M': 3, 'L': 4, 'XL': 5}


def szgroup(cat, size):
    """サイズの系統。ソックスの M はメンズの M41 とは別物なので先に分ける"""
    if cat == 'socks':
        return 'sock'
    if size.startswith('W'):
        return 'W'
    if size.startswith('M'):
        return 'M'
    return 'tabi'


def szkey(s):
    if s in SOCKORD:
        return (1, SOCKORD[s])
    try:
        return (0, float(s.lstrip('WM').replace('cm', '')))
    except ValueError:
        return (2, 999.0)


# ---------------------------------------------------------------- 集計
J = json.load(io.open(os.path.join(HERE, 'agg.json'), encoding='utf-8'))
SALES = [o for o in J['orders'] if o['bucket'] == 'sales']
PERIOD = J['meta']['period']

sku = collections.defaultdict(lambda: dict(units=0, list=0.0, net=0.0, off=0, on=0, wh=0,
                                           first='9999-99-99', last='0000-00-00',
                                           cat='', brand=''))
for o in SALES:
    seg = 'off' if o['seg'].startswith(u'オフライン') else ('wh' if o['seg'] == u'卸売' else 'on')
    prod = [it for it in o['items'] if it['cat'] not in ('test', 'fee')]
    base = sum(it['total'] for it in prod)
    # 送料などの商品でない行を注文合計から外し、残りを明細金額の比で割り振る
    nonprod = sum(it['total'] for it in o['items'] if it['cat'] in ('test', 'fee'))
    share = (o['total'] - nonprod) / base if base else 0.0
    for it in prod:
        r = sku[(it['brand'], it['model'], it['color'] or u'(記録なし)', it['size'] or u'(記録なし)')]
        r['units'] += it['qty']
        r['list'] += it['total']
        r['net'] += it['total'] * share
        r[seg] += it['qty']
        r['cat'] = it['cat']
        r['brand'] = it['brand']
        r['first'] = min(r['first'], o['date'])
        r['last'] = max(r['last'], o['date'])

ROWS = [dict(v, model=k[1], color=k[2], size=k[3], grp=szgroup(v['cat'], k[3]))
        for k, v in sku.items()]
TOT_U = sum(r['units'] for r in ROWS)
TOT_N = sum(r['net'] for r in ROWS)
TOT_L = sum(r['list'] for r in ROWS)
SUM_OFF = sum(r['off'] for r in ROWS)
SUM_ON = sum(r['on'] for r in ROWS)
SUM_WH = sum(r['wh'] for r in ROWS)


def roll(keyf):
    d = collections.defaultdict(lambda: dict(units=0, list=0.0, net=0.0, off=0, on=0, wh=0,
                                             first='9999-99-99', last='0000-00-00',
                                             cat='', brand=''))
    for r in ROWS:
        v = d[keyf(r)]
        for f in ('units', 'list', 'net', 'off', 'on', 'wh'):
            v[f] += r[f]
        v['first'] = min(v['first'], r['first'])
        v['last'] = max(v['last'], r['last'])
        v['cat'] = r['cat']
        v['brand'] = r['brand']
    return d


BYMODEL = roll(lambda r: r['model'])
BYCOLOR = roll(lambda r: (r['model'], r['color']))
MORD = {m: -v['units'] for m, v in BYMODEL.items()}   # 足数の多い順


def D(s):
    return datetime.date(*map(int, s.split('-')))


wb = Workbook()


def header(ws, labels, widths, row=1, freeze=True, autofilter=True):
    for i, (lab, w) in enumerate(zip(labels, widths), 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    if freeze:
        ws.freeze_panes = ws.cell(row=row + 1, column=1)
    if autofilter:
        ws.auto_filter.ref = u'A%d:%s%d' % (row, get_column_letter(len(labels)), row)


def put(ws, r, vals, fmts, total=False):
    for i, (v, f) in enumerate(zip(vals, fmts), 1):
        c = ws.cell(row=r, column=i, value=v)
        if f:
            c.number_format = f
        if total:
            c.fill = TOT_FILL
            c.font = Font(bold=True)
            c.border = UNDER
    return r + 1


# ---------------------------------------------------------------- 1. 概要
ws = wb.active
ws.title = u'概要'
ws['A1'] = u'Barefootinc Malaysia 商品別販売実績'
ws['A1'].font = Font(bold=True, size=14)
INFO = [(u'対象期間', u'%s 〜 %s' % tuple(PERIOD)),
        (u'純売上(商品分)', round(TOT_N, 2)),
        (u'販売足数', TOT_U),
        (u'注文数', len(SALES)),
        (u'SKU数(モデル×カラー×サイズ)', len(ROWS)),
        (u'売上の基準', u'「売上(実売)」= 注文合計を明細に按分した実売ベース。'
                      u'「売上(定価)」= 明細の定価合計で、値引きが引かれていない')]
for i, (k, v) in enumerate(INFO, 3):
    ws.cell(row=i, column=1, value=k).font = Font(bold=True)
    c = ws.cell(row=i, column=2, value=v)
    if isinstance(v, float):
        c.number_format = MONEY
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 24

HR = 10
header(ws, [u'モデル', u'ブランド', u'区分', u'足数', u'足数\n構成比', u'売上(実売)',
            u'売上\n構成比', u'平均\n実売単価', u'売上(定価)', u'店舗\n(足)',
            u'オンライン\n(足)', u'卸売\n(足)', u'初回販売日', u'最終販売日'],
       [22, 20, 10, 8, 9, 13, 9, 11, 13, 9, 11, 9, 13, 13], row=HR)
FM = [None, None, None, INTF, PCTF, MONEY, PCTF, MONEY, MONEY, INTF, INTF, INTF, DATEF, DATEF]
r = HR + 1
for m, v in sorted(BYMODEL.items(), key=lambda x: (MORD[x[0]], x[0])):
    r = put(ws, r, [m, v['brand'], CATNAME.get(v['cat'], v['cat']), v['units'],
                    v['units'] / float(TOT_U), round(v['net'], 2), v['net'] / TOT_N,
                    round(v['net'] / v['units'], 2), round(v['list'], 2),
                    v['off'], v['on'], v['wh'], D(v['first']), D(v['last'])], FM)
put(ws, r, [u'合計', None, None, TOT_U, 1.0, round(TOT_N, 2), 1.0, round(TOT_N / TOT_U, 2),
            round(TOT_L, 2), SUM_OFF, SUM_ON, SUM_WH, None, None], FM, True)

# ---------------------------------------------------------------- 2. モデル×カラー
ws = wb.create_sheet(u'モデル×カラー')
header(ws, [u'モデル', u'カラー', u'区分', u'足数', u'構成比', u'売上(実売)',
            u'平均実売単価', u'店舗(足)', u'オンライン(足)', u'卸売(足)',
            u'サイズ内訳', u'初回販売日', u'最終販売日'],
       [22, 20, 10, 8, 9, 13, 12, 9, 11, 9, 46, 13, 13])
FM = [None, None, None, INTF, PCTF, MONEY, MONEY, INTF, INTF, INTF, None, DATEF, DATEF]
szof = collections.defaultdict(collections.Counter)
for x in ROWS:
    szof[(x['model'], x['color'])][x['size']] += x['units']
r = 2
for k, v in sorted(BYCOLOR.items(), key=lambda x: (MORD[x[0][0]], x[0][0], -x[1]['units'])):
    szs = u' '.join(u'%s×%d' % (s, n)
                    for s, n in sorted(szof[k].items(), key=lambda x: szkey(x[0])))
    r = put(ws, r, [k[0], k[1], CATNAME.get(v['cat'], v['cat']), v['units'],
                    v['units'] / float(TOT_U), round(v['net'], 2),
                    round(v['net'] / v['units'], 2), v['off'], v['on'], v['wh'],
                    szs, D(v['first']), D(v['last'])], FM)
put(ws, r, [u'合計', None, None, TOT_U, 1.0, round(TOT_N, 2), round(TOT_N / TOT_U, 2),
            SUM_OFF, SUM_ON, SUM_WH, None, None, None], FM, True)

# ---------------------------------------------------------------- 3. SKU明細
ws = wb.create_sheet(u'SKU明細')
header(ws, [u'ブランド', u'モデル', u'区分', u'サイズ区分', u'カラー', u'サイズ', u'足数',
            u'売上(実売)', u'売上(定価)', u'平均実売単価', u'店舗(足)',
            u'オンライン(足)', u'卸売(足)', u'初回販売日', u'最終販売日'],
       [20, 22, 10, 17, 20, 10, 8, 13, 13, 12, 9, 11, 9, 13, 13])
FM = [None, None, None, None, None, None, INTF, MONEY, MONEY, MONEY,
      INTF, INTF, INTF, DATEF, DATEF]
ROWS.sort(key=lambda x: (MORD[x['model']], x['model'], x['color'], szkey(x['size'])))
r = 2
for x in ROWS:
    r = put(ws, r, [x['brand'], x['model'], CATNAME.get(x['cat'], x['cat']),
                    GRPNAME[x['grp']], x['color'], x['size'], x['units'],
                    round(x['net'], 2), round(x['list'], 2),
                    round(x['net'] / x['units'], 2), x['off'], x['on'], x['wh'],
                    D(x['first']), D(x['last'])], FM)
put(ws, r, [u'合計', None, None, None, None, None, TOT_U, round(TOT_N, 2), round(TOT_L, 2),
            round(TOT_N / TOT_U, 2), SUM_OFF, SUM_ON, SUM_WH, None, None], FM, True)
ws.auto_filter.ref = u'A1:O%d' % (r - 1)

# ---------------------------------------------------------------- 4. サイズ別マトリクス
ws = wb.create_sheet(u'サイズ別マトリクス')
ws['A1'] = u'サイズ別 販売足数(モデル×カラー)'
ws['A1'].font = Font(bold=True, size=14)
ws['A2'] = u'色が濃いほど売れているサイズ。空欄は販売なし。'
ws['A2'].font = Font(size=10, color='767676')
ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 20
ws.freeze_panes = 'C4'

r = 4
for g in GRPORD:
    sub = [x for x in ROWS if x['grp'] == g]
    if not sub:
        continue
    cols = sorted({x['size'] for x in sub}, key=szkey)
    last = len(cols) + 3                       # 合計列
    ws.cell(row=r, column=1,
            value=u'%s  計%d足' % (GRPNAME[g], sum(x['units'] for x in sub))
            ).font = Font(bold=True, size=12)
    r += 1
    for i, lab in enumerate([u'モデル', u'カラー'] + cols + [u'合計'], 1):
        c = ws.cell(row=r, column=i, value=lab)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal='center')
        if i > 2:
            ws.column_dimensions[get_column_letter(i)].width = 9
    r += 1
    grid = {}
    for x in sub:
        grid[(x['model'], x['color'], x['size'])] = x['units']
    keys = sorted({(x['model'], x['color']) for x in sub},
                  key=lambda k: (MORD[k[0]], k[0], k[1]))
    first = r
    for k in keys:
        ws.cell(row=r, column=1, value=k[0])
        ws.cell(row=r, column=2, value=k[1])
        for ci, s in enumerate(cols, 3):
            n = grid.get((k[0], k[1], s))
            if n:
                ws.cell(row=r, column=ci, value=n).alignment = Alignment(horizontal='center')
        c = ws.cell(row=r, column=last, value=sum(grid.get((k[0], k[1], s), 0) for s in cols))
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal='center')
        r += 1
    ws.cell(row=r, column=1, value=u'合計')
    for ci, s in enumerate(cols, 3):
        ws.cell(row=r, column=ci, value=sum(grid.get((k[0], k[1], s), 0) for k in keys))
    ws.cell(row=r, column=last, value=sum(x['units'] for x in sub))
    for ci in range(1, last + 1):
        c = ws.cell(row=r, column=ci)
        c.fill = TOT_FILL
        c.font = Font(bold=True)
        c.border = UNDER
        if ci > 2:
            c.alignment = Alignment(horizontal='center')
    ws.conditional_formatting.add(
        u'C%d:%s%d' % (first, get_column_letter(last - 1), r - 1),
        ColorScaleRule(start_type='num', start_value=1, start_color='FFF6E7DA',
                       end_type='max', end_color='FFD2691E'))
    r += 2

wb.save(OUT)
print('wrote %s  (%d SKU / %d足 / RM %.2f)' % (OUT, len(ROWS), TOT_U, TOT_N))
